# -*- coding: utf-8 -*-
"""
Script para actualizar el documento oficial del PPI: AburráSegura.docx
- Poblado exhaustivo de la Tabla 1 de Abreviaturas técnicas (ordenadas alfabéticamente).
- Sustitución del texto instructivo de borrador en el párrafo 71 por redacción formal.
- Completado de la columna 'Fase / entregable asociado' en la Tabla 2 de Objetivos Específicos.
"""

import sys
import copy
from pathlib import Path
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def update_document():
    doc_path = Path("AburráSegura.docx").resolve()
    print(f"Abriendo documento: {doc_path}")
    
    doc = docx.Document(str(doc_path))
    
    # 1. Actualizar párrafo 71 (reemplazo de marcador de borrador)
    p71 = doc.paragraphs[71]
    p71.text = "A continuación, se relacionan las principales siglas, acrónimos y conceptos técnicos empleados en el presente documento:"
    p71.runs[0].font.size = Pt(10.5)
    p71.runs[0].italic = True
    
    # 2. Lista completa y rigurosa de abreviaturas del proyecto
    abbreviations = [
        ("API", "Application Programming Interface (Interfaz de Programación de Aplicaciones; conjunto de protocolos y definiciones para la comunicación e integración entre componentes de software)."),
        ("B-Tree", "Árbol B (Estructura de datos balanceada y autoequilibrada empleada en sistemas de bases de datos para optimizar búsquedas y consultas indexadas)."),
        ("CORS", "Cross-Origin Resource Sharing (Intercambio de Recursos de Origen Cruzado; estándar de seguridad web para habilitar peticiones HTTP entre diferentes orígenes)."),
        ("CPU", "Central Processing Unit (Unidad Central de Procesamiento)."),
        ("CSV", "Comma-Separated Values (Valores Separados por Comas; formato estándar de archivo plano para el intercambio de datos tabulares)."),
        ("DAG", "Directed Acyclic Graph (Grafo Dirigido Acíclico)."),
        ("DBSCAN", "Density-Based Spatial Clustering of Applications with Noise (Algoritmo de aprendizaje no supervisado para agrupamiento espacial basado en densidad)."),
        ("ETL", "Extract, Transform, Load (Extracción, Transformación y Carga; canalización computacional de depuración, normalización y almacenamiento de datos masivos)."),
        ("GeoJSON", "Geographic JavaScript Object Notation (Estándar abierto para la serialización de entidades y geometrías espaciales mediante formato JSON)."),
        ("GPS", "Global Positioning System (Sistema de Posicionamiento Global)."),
        ("HTTP / HTTPS", "Hypertext Transfer Protocol / Secure (Protocolo Seguro de Transferencia de Hipertexto para la transmisión de información en la web)."),
        ("JSON", "JavaScript Object Notation (Notación de Objetos de JavaScript; estándar abierto y liviano para serialización e intercambio de datos estructurados)."),
        ("KD-Tree", "k-dimensional Tree (Árbol k-dimensional; estructura de datos de partición del espacio para acelerar búsquedas de puntos vecinos más cercanos en coordenadas geográficas)."),
        ("MEData", "Portal Oficial de Datos Abiertos de la Alcaldía de Medellín (Fuente pública institucional de siniestralidad vial)."),
        ("Min-Heap", "Montículo Binario de Mínimos (Estructura de datos para cola de prioridad donde la raíz almacena el elemento de menor peso o coste acumulado)."),
        ("OSM", "OpenStreetMap (Plataforma colaborativa y base de datos geográfica abierta a nivel global)."),
        ("OSRM", "Open Source Routing Machine (Motor de enrutamiento vial de alto rendimiento en C++ para cálculo de rutas sobre redes OpenStreetMap)."),
        ("PPI", "Proyecto Pedagógico Integrador (Estrategia formativa institucional de la Facultad de Ingeniería del Tecnológico de Antioquia)."),
        ("RAM", "Random Access Memory (Memoria de Acceso Aleatorio de trabajo del sistema)."),
        ("REST", "Representational State Transfer (Transferencia de Estado Representacional; estilo arquitectónico para servicios web desacoplados y sin estado)."),
        ("SPA", "Single Page Application (Aplicación Web de Página Única con carga dinámica de componentes)."),
        ("SQLite", "Motor de base de datos relacional SQL embebido, transaccional, autónomo y sin dependencias de servidor."),
        ("TdeA", "Tecnológico de Antioquia Institución Universitaria (Medellín, Colombia)."),
        ("UI / UX", "User Interface / User Experience (Interfaz de Usuario y Experiencia de Usuario)."),
        ("WGS84", "World Geodetic System 1984 (Sistema geodésico estándar mundial de referencia cartográfica para coordenadas de latitud y longitud).")
    ]
    
    t1 = doc.tables[1]
    print(f"Table 1 original rows: {len(t1.rows)}")
    
    # Asegurar que Table 1 tenga 1 fila de encabezado + 25 filas de abreviaturas
    while len(t1.rows) - 1 < len(abbreviations):
        new_tr = copy.deepcopy(t1.rows[1]._tr)
        t1._tbl.append(new_tr)
        
    print(f"Table 1 expanded rows: {len(t1.rows)}")
    
    # Poblar cada fila de la Tabla de Abreviaturas
    for idx, (sigla, significado) in enumerate(abbreviations):
        row = t1.rows[idx + 1]
        
        # Celda 0: Sigla (Centrada / Negrita)
        c0 = row.cells[0]
        c0.text = ""
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(sigla)
        r0.bold = True
        r0.font.size = Pt(9.5)
        
        # Celda 1: Significado (Justificado / Normal)
        c1 = row.cells[1]
        c1.text = ""
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r1 = p1.add_run(significado)
        r1.font.size = Pt(9.5)
        
    print("Tabla 1 (Abreviaturas) poblada con éxito.")
    
    # 3. Poblar Tabla 2 (Objetivos Específicos: Columna 'Fase / entregable asociado')
    t2 = doc.tables[2]
    
    deliverables = [
        {
            "row_idx": 1,
            "phase": "Fase 1: Análisis y Procesamiento de Datos\n",
            "bullets": (
                "• Pipeline ETL geoespacial de depuración de 251.150 siniestros MEData con validación WGS84.\n"
                "• Base de datos relacional SQLite indexada (data/processed/medellin_incidents.db).\n"
                "• Matriz de delimitación de puntos críticos viales (hotspots) y cuadrícula analítica de riesgo."
            )
        },
        {
            "row_idx": 2,
            "phase": "Fase 2: Diseño e Implementación Algorítmica\n",
            "bullets": (
                "• Modelo topológico de la red vial en grafo dirigido ponderado G = (V, E) serializado (data/processed/medellin_graph.json).\n"
                "• Algoritmo de Dijkstra con cola de prioridad min-heap (O((V + E) log V)) y función de coste multiobjetivo con aversión al riesgo.\n"
                "• API REST en FastAPI (/api/v1/routing/calculate) con trazado vial continuo OSRM de alta resolución tipo Waze."
            )
        },
        {
            "row_idx": 3,
            "phase": "Fase 3: Pruebas, Benchmarking e Integración\n",
            "bullets": (
                "• Laboratorio de benchmarking asintótico en tiempo real (O(1) a O(n²)) evaluado sobre MEData (/api/v1/benchmarks/complexity).\n"
                "• Suite automatizada de 44 pruebas unitarias y de integración en Pytest.\n"
                "• Aplicación web interactiva SPA (React 18 + Leaflet) con cálculo comparativo de rutas en vivo y panel de métricas de seguridad."
            )
        }
    ]
    
    for item in deliverables:
        row = t2.rows[item["row_idx"]]
        c2 = row.cells[2]
        c2.text = ""
        p = c2.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        # Fase (Negrita)
        r_phase = p.add_run(item["phase"])
        r_phase.bold = True
        r_phase.font.size = Pt(9.5)
        
        # Etiqueta 'Entregables asociados:' (Cursiva)
        r_hdr = p.add_run("Entregables asociados:\n")
        r_hdr.italic = True
        r_hdr.font.size = Pt(9.0)
        
        # Viñetas de entregables
        r_b = p.add_run(item["bullets"])
        r_b.font.size = Pt(9.0)
        
    print("Tabla 2 (Objetivos Específicos) actualizada con éxito.")
    
    # Guardar documento en la ruta oficial
    doc.save(str(doc_path))
    print(f"Documento guardado exitosamente en: {doc_path}")

if __name__ == "__main__":
    update_document()
